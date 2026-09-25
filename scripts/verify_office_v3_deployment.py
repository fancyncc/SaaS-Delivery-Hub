"""Exercise PPTX/XLSX upload, indexing, and location-correct V3 evidence."""

import io
import json
import time
from pathlib import Path
from uuid import uuid4

import httpx
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from verify_index_lifecycle import cleanup

ROOT = Path(__file__).resolve().parents[1]


def sample_pptx():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Project audit code JUPITER 82351"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def sample_xlsx():
    workbook = Workbook()
    workbook.active.title = "Project audit"
    workbook.active["A1"] = "Audit code"
    workbook.active["B2"] = "SATURN 61422"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def main():
    workspace = ROOT / "data" / "demo_workspace"
    admin = json.loads((workspace / "credentials.json").read_text(encoding="utf-8"))["admin"]
    manifest = json.loads((workspace / "manifest.json").read_text(encoding="utf-8"))
    uploaded = []
    with httpx.Client(base_url="http://127.0.0.1:8080", trust_env=False, timeout=45) as client:
        def post(path, payload):
            response = client.post(path, json=payload,
                headers={"X-CSRF-Token": client.cookies.get("saas_csrf", "")})
            response.raise_for_status()
            return response.json()["data"]

        post("/api/auth/login", {"username": admin["username"], "password": admin["password"]})
        post(f"/api/auth/spaces/{manifest['tenant_id']}/switch", {})
        try:
            for filename, raw, code, location in (
                ("audit.pptx", sample_pptx(), "JUPITER 82351", {"slide": 1}),
                ("audit.xlsx", sample_xlsx(), "SATURN 61422", {"sheet": "Project audit", "row": 2}),
            ):
                title = f"索引验收临时资料-{uuid4().hex[:12]}"
                response = client.post("/api/knowledge/files",
                    data={"title": title, "version": "1", "module": "Office 验收", "license": "本机验收"},
                    files={"file": (filename, raw)},
                    headers={"X-CSRF-Token": client.cookies.get("saas_csrf", "")})
                response.raise_for_status()
                document_id = response.json()["data"]["id"]
                uploaded.append((document_id, title))
                for _ in range(60):
                    detail = client.get(f"/api/knowledge/{document_id}").json()["data"]
                    if detail["v3"]["phase"] in {"ready", "failed"}:
                        break
                    time.sleep(3)
                assert detail["v3"]["phase"] == "ready", detail["v3"]
                result = post("/api/knowledge/inspect", {"question": f"What is the audit code {code}?"})
                matches = [chunk for chunk in result["chunks"] if chunk["origin_id"] == document_id]
                assert matches and all(all(chunk["location"].get(k) == v for k, v in location.items())
                                       for chunk in matches), result
                assert code in result["evidence_text"], result
                print(json.dumps({"format": filename.rsplit(".", 1)[-1], "v3": "passed",
                                  "location": location}, ensure_ascii=True))
        finally:
            for document_id, title in uploaded:
                cleanup(document_id, title)


if __name__ == "__main__":
    main()

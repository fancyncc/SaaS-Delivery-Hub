"""Generate reproducible binary boundary fixtures for V3's offline evaluation."""

import json
from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

ROOT = Path(__file__).resolve().parents[1] / "evaluations" / "v3"
DOMAINS = ("inventory", "education", "facilities", "publishing", "events", "software")
FORMATS = ("pdf", "pptx", "xlsx")


def source_fields(domain: str) -> dict[str, str]:
    lines = (ROOT / "fixtures" / f"{domain}.md").read_text(encoding="utf-8").splitlines()
    values = dict(line.split(": ", 1) for line in lines if ": " in line)
    values["conditions"] = values["conditions"].split("。 ", 1)[1]
    return values


def write_pdf(path: Path, fields: dict[str, str]) -> None:
    writer = PdfWriter()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    font_ref = writer._add_object(font)
    for lines in (
        (f"id: {fields['id']}", f"limit: {fields['limit']}", f"fields: {fields['fields']}"),
        (f"conditions: {fields['conditions']}", fields["distractor"].split(". ", 1)[0] + "."),
    ):
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})
        })
        content = ["BT /F1 12 Tf 40 740 Td"]
        for line in lines:
            safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            content.append(f"({safe}) Tj 0 -22 Td")
        content.append("ET")
        stream = DecodedStreamObject()
        stream.set_data("\n".join(content).encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as output:
        writer.write(output)


def write_pptx(path: Path, fields: dict[str, str]) -> None:
    presentation = Presentation()
    for lines in (
        (f"id: {fields['id']}", f"limit: {fields['limit']}", f"fields: {fields['fields']}"),
        (f"conditions: {fields['conditions']}", fields["distractor"]),
    ):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(8), Inches(4))
        box.text = "\n".join(lines)
    presentation.save(path)


def write_xlsx(path: Path, fields: dict[str, str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Policy"
    for key in ("id", "limit", "fields", "conditions", "distractor"):
        sheet.append((key, fields[key]))
    workbook.save(path)


def cases_for(domain: str, fmt: str, fields: dict[str, str]) -> list[dict]:
    identifier = fields["id"]
    split = "calibration" if domain in DOMAINS[:3] else "validation"
    questions = (
        ("fact", f"What is the limit for {identifier}?", [fields["limit"]], ["limit"]),
        ("enumeration", f"List all fields of {identifier}.",
         ["id", "owner", "created_at", "status"], ["fields"]),
        ("conditions", f"What conditions and exceptions apply to {identifier}?",
         [fields["conditions"]], ["conditions"]),
        ("noanswer", f"What is the insurance premium of {identifier}?", [], []),
    )
    return [{
        "id": f"{domain}-{fmt}-{index}",
        "split": split,
        "domain": domain,
        "document_group": domain,
        "format": fmt,
        "language": "en",
        "kind": kind,
        "question": question,
        "document": f"fixtures/{domain}.{fmt}",
        "necessary_facts": facts,
        "source_keys": keys,
        "protected_conditions": [fields["conditions"]] if kind == "conditions" else [],
        "distractors": ["ARCHIVE-999", "999"],
        "allow_no_answer": kind == "noanswer",
        "provenance": "synthetic_binary_fixture",
        "review_status": "pending_independent_review",
    } for index, (kind, question, facts, keys) in enumerate(questions)]


def main() -> None:
    manifest = ROOT / "cases.jsonl"
    existing = [json.loads(line) for line in manifest.read_text(encoding="utf-8").splitlines()]
    cases = [case for case in existing if case["format"] not in FORMATS]
    for domain in DOMAINS:
        fields = source_fields(domain)
        for fmt, writer in (("pdf", write_pdf), ("pptx", write_pptx), ("xlsx", write_xlsx)):
            writer(ROOT / "fixtures" / f"{domain}.{fmt}", fields)
            cases.extend(cases_for(domain, fmt, fields))
    manifest.write_text(
        "".join(json.dumps(case, ensure_ascii=False) + "\n" for case in cases),
        encoding="utf-8",
    )
    print(f"Wrote {len(cases)} V3 boundary cases across {len(set(c['format'] for c in cases))} formats")


if __name__ == "__main__":
    main()

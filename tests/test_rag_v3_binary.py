import io
import json
import subprocess
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import select

from backend.config import get_settings
from backend.db import SessionLocal
from backend.models import KnowledgeDocument
from backend.rag_v3 import visible_documents
from backend.rag_v3_index import index_document
from backend.rag_v3_models import V3Document, V3Unit
from backend.rag_v3_parse import parse
from backend.rag_v3_parse_runner import parse_isolated


def pdf_with_text(*texts):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=400, height=400)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
            DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 300 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_pdf_pages_are_distinct_and_isolated():
    raw = pdf_with_text("First page limit is 80.", "Second page limit is 90.")
    parsed = parse("limits.pdf", raw)
    assert [node.location["page"] for node in parsed.nodes] == [1, 2]
    assert "80" in parsed.nodes[0].text and "90" not in parsed.nodes[0].text
    assert "90" in parsed.nodes[1].text and "80" not in parsed.nodes[1].text
    isolated = parse_isolated("limits.pdf", raw)
    assert isolated.dump() == parsed.dump()


def test_pdf_scanned_page_is_not_silently_omitted(monkeypatch):
    image = Image.new("RGB", (100, 100), "black")
    buffer = io.BytesIO()
    image.save(buffer, format="PDF")
    monkeypatch.setattr("backend.rag_v3_binary._ocr_page", lambda raw, page: "OCR result")
    parsed = parse("scan.pdf", buffer.getvalue())
    assert parsed.nodes[0].text == "OCR result"
    assert parsed.nodes[0].location["page"] == 1
    assert "OCR" in parsed.warnings[0]


def test_blank_pdf_pages_do_not_consume_ocr_limit(monkeypatch):
    writer = PdfWriter()
    from pypdf import PdfReader

    writer.add_page(PdfReader(io.BytesIO(pdf_with_text("Searchable policy text"))).pages[0])
    for _ in range(31):
        writer.add_blank_page(width=400, height=400)
    stream = io.BytesIO()
    writer.write(stream)
    monkeypatch.setattr("backend.rag_v3_binary._ocr_page", lambda *_: pytest.fail("Blank page reached OCR"))
    parsed = parse("policy.pdf", stream.getvalue())
    assert [node.location["page"] for node in parsed.nodes] == [1]
    assert parsed.ocr_pages == 0


def test_unreadable_scanned_page_fails_instead_of_partial_index(monkeypatch):
    image = Image.new("RGB", (100, 100), "black")
    output = io.BytesIO()
    image.save(output, format="PDF")
    monkeypatch.setattr("backend.rag_v3_binary._ocr_page", lambda raw, page: "")
    with pytest.raises(HTTPException) as exc:
        parse("unreadable.pdf", output.getvalue())
    assert exc.value.status_code == 422 and "第 1 页" in str(exc.value.detail)


def test_pdf_encrypted_and_bad_magic():
    with pytest.raises(HTTPException) as exc:
        parse("x.pdf", b"not a PDF")
    assert exc.value.status_code == 422
    writer = PdfWriter()
    writer.add_blank_page(width=400, height=400)
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(HTTPException) as exc:
        parse("secret.pdf", buffer.getvalue())
    assert "加密" in str(exc.value.detail)


def test_parser_timeout_is_reported(monkeypatch):
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("parser", 1)

    monkeypatch.setattr("backend.rag_v3_parse_runner.subprocess.run", timeout)
    with pytest.raises(HTTPException) as exc:
        parse_isolated("slow.pdf", b"%PDF-")
    assert exc.value.status_code == 422 and "超时" in str(exc.value.detail)


def test_pptx_slide_and_xlsx_cell_locations():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    textbox.text = "Project launch checklist"
    stream = io.BytesIO()
    presentation.save(stream)
    parsed = parse("slides.pptx", stream.getvalue())
    assert any(node.location["slide"] == 1 and "checklist" in node.text for node in parsed.nodes)

    workbook = Workbook()
    workbook.active.title = "Budget"
    workbook.active["A1"] = "Item"
    workbook.active["B2"] = 42
    stream = io.BytesIO()
    workbook.save(stream)
    parsed = parse("budget.xlsx", stream.getvalue())
    assert any(node.location == {"sheet": "Budget", "row": 2, "column": 2}
               and node.text == "42" for node in parsed.nodes)


def test_binary_evaluation_fixtures_preserve_annotated_facts():
    root = Path("evaluations/v3")
    cases = [json.loads(line) for line in (root / "cases.jsonl").read_text(encoding="utf-8").splitlines()]
    for case in cases:
        if case["format"] not in {"pdf", "pptx", "xlsx"}:
            continue
        path = root / case["document"]
        parsed = parse(path.name, path.read_bytes())
        text = "\n".join(node.text for node in parsed.nodes)
        assert all(fact in text for fact in case["necessary_facts"]), case["id"]
        expected = {"pdf": "page", "pptx": "slide", "xlsx": "sheet"}[case["format"]]
        assert any(expected in node.location for node in parsed.nodes), case["id"]


async def test_pdf_upload_registers_original_for_v3(client, monkeypatch):
    raw = pdf_with_text("Project policy limit is 80.", "Second page limit is 90.")
    response = await client.post(
        "/api/knowledge/files",
        data={"title": "PDF policy", "version": "1", "module": "test", "license": "internal"},
        files={"file": ("policy.pdf", raw, "application/pdf")},
    )
    assert response.status_code == 200, response.text
    document_id = response.json()["data"]["id"]
    async with SessionLocal() as session:
        source = await session.scalar(select(V3Document).where(V3Document.origin_id == document_id))
        document = await session.get(KnowledgeDocument, document_id)
        assert source.phase == "pending" and source.filename == "policy.pdf" and source.raw == raw
        assert document.index_status == "v3_pending"
        async def counts(texts):
            return [len(text.split()) for text in texts]
        monkeypatch.setattr("backend.rag_v3_index.model_counts", counts)
        await index_document(session, source)
        await session.flush()
        units = (await session.scalars(select(V3Unit).where(V3Unit.document_id == source.id)
                                       .order_by(V3Unit.ordinal))).all()
        assert source.phase == "ready" and [unit.location["page"] for unit in units] == [1, 2]
        assert units[0].next_id is None and units[1].previous_id is None
        assert "## 第 1 页" in document.body and "## 第 2 页" in document.body
        monkeypatch.setattr(get_settings(), "rag_pdf_enabled", False)
        hidden = (await session.execute(visible_documents(document.tenant_id, []))).all()
        assert source.id not in {row[0].id for row in hidden}

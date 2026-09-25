"""Bounded, location-preserving extraction for binary knowledge documents."""

import io
import re
import zipfile

from fastapi import HTTPException

PDF_MAX_PAGES = 100
PDF_MAX_OCR_PAGES = 30
OFFICE_MAX_UNCOMPRESSED = 40 * 1024 * 1024


def validate_binary(suffix: str, raw: bytes) -> None:
    if suffix == ".pdf":
        if not raw.startswith(b"%PDF-"):
            raise HTTPException(422, "文件内容不是有效 PDF")
        return
    if suffix not in {".docx", ".pptx", ".xlsx"}:
        return
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            names = archive.namelist()
            if len(names) > 10000 or sum(info.file_size for info in archive.infolist()) > OFFICE_MAX_UNCOMPRESSED:
                raise HTTPException(413, "Office 文件解压后内容过大")
            marker = {".docx": "word/document.xml", ".pptx": "ppt/presentation.xml", ".xlsx": "xl/workbook.xml"}[suffix]
            if marker not in names:
                raise HTTPException(422, "文件内容与扩展名不匹配")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise HTTPException(415, "不支持包含宏的 Office 文件")
    except zipfile.BadZipFile:
        raise HTTPException(422, "Office 文件损坏或格式无效") from None


def _ocr_page(raw: bytes, index: int) -> str:
    import pypdfium2 as pdfium
    import pytesseract

    document = pdfium.PdfDocument(raw)
    try:
        page = document[index]
        try:
            if page.get_width() * page.get_height() * 4 > 25_000_000:
                raise HTTPException(413, "PDF 页面图像过大")
            image = page.render(scale=2).to_pil()
            try:
                try:
                    return pytesseract.image_to_string(image, lang="chi_sim+eng", timeout=20).strip()
                except pytesseract.TesseractNotFoundError:
                    raise HTTPException(503, "OCR 服务未安装或不可用") from None
                except pytesseract.TesseractError as exc:
                    if "Error opening data file" in str(exc):
                        raise HTTPException(503, "OCR 语言包不可用") from None
                    raise HTTPException(422, "PDF 页面 OCR 失败，请检查扫描质量") from None
            finally:
                image.close()
        finally:
            page.close()
    finally:
        document.close()


def parse_pdf(raw: bytes, add, result) -> None:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw), strict=True)
    if reader.is_encrypted:
        raise HTTPException(422, "PDF 已加密，请解密后上传")
    if len(reader.pages) > PDF_MAX_PAGES:
        raise HTTPException(413, f"PDF 不能超过 {PDF_MAX_PAGES} 页")
    ocr_pages = 0
    for page_number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        has_images = bool(page.images)
        if not text and not has_images and not page.get("/Annots"):
            contents = page.get_contents()
            if contents is None or not contents.get_data().strip():
                continue
        needs_ocr = not text or (has_images and len(text) < 80)
        if needs_ocr:
            ocr_pages += 1
            result.ocr_pages += 1
            if ocr_pages > PDF_MAX_OCR_PAGES:
                raise HTTPException(413, f"需要 OCR 的页面不能超过 {PDF_MAX_OCR_PAGES} 页")
            recognized = _ocr_page(raw, page_number - 1)
            if has_images and not recognized and not text:
                raise HTTPException(422, f"PDF 第 {page_number} 页无法识别文字")
            if recognized and re.sub(r"\s+", "", recognized) not in re.sub(r"\s+", "", text):
                text = (text + "\n" if text else "") + recognized
            if recognized:
                result.warnings.append(f"第 {page_number} 页使用 OCR，识别内容请人工核对")
        elif has_images:
            result.warnings.append(f"第 {page_number} 页仅提取文字层，图片内容未识别")
        for block_index, block in enumerate(re.split(r"\n\s*\n", text), 1):
            if block.strip():
                add("paragraph", block.strip(), {"page": page_number, "block": block_index})
    if not result.nodes:
        raise HTTPException(422, "PDF 没有可检索文字")


def parse_pptx(raw: bytes, add, result) -> None:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(raw))
    if len(presentation.slides) > 200:
        raise HTTPException(413, "演示文稿不能超过 200 页")
    for slide_number, slide in enumerate(presentation.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape else ""
        parent = add("heading", title, {"slide": slide_number}, heading=title) if title else None
        for shape_number, shape in enumerate(slide.shapes, 1):
            if shape.has_table:
                for row_number, row in enumerate(shape.table.rows, 1):
                    for column_number, cell in enumerate(row.cells, 1):
                        if cell.text.strip():
                            add("table_cell", cell.text.strip(),
                                {"slide": slide_number, "shape": shape_number, "row": row_number, "column": column_number},
                                record=f"slide:{slide_number}:table:{shape_number}:row:{row_number}",
                                context=f"列 {column_number}", parent=parent, heading=title)
            elif shape.has_text_frame and (title_shape is None or shape.shape_id != title_shape.shape_id):
                for paragraph_number, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                    if paragraph.text.strip():
                        add("paragraph", paragraph.text.strip(),
                            {"slide": slide_number, "shape": shape_number, "paragraph": paragraph_number},
                            parent=parent, heading=title)
            elif shape.shape_type == 13:
                result.warnings.append(f"第 {slide_number} 页图片未识别")


def parse_xlsx(raw: bytes, add, result) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    try:
        if len(workbook.worksheets) > 50:
            raise HTTPException(413, "工作簿不能超过 50 个工作表")
        for sheet in workbook.worksheets:
            if sheet.max_row > 10000 or sheet.max_column > 200:
                raise HTTPException(413, f"工作表 {sheet.title} 超过行列限制")
            parent = add("heading", sheet.title, {"sheet": sheet.title}, heading=sheet.title)
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                for column_number, value in enumerate(row, 1):
                    if value is not None and str(value).strip():
                        add("table_cell", str(value),
                            {"sheet": sheet.title, "row": row_number, "column": column_number},
                            record=f"sheet:{sheet.title}:row:{row_number}",
                            context=f"列 {column_number}", parent=parent, heading=sheet.title)
    finally:
        workbook.close()
    result.warnings.append("电子表格只读取已保存的单元格值，不计算公式")
